# python-docx 라이브러리를 사용해 K-food 관련 자료를 워드 문서(k-food.docx)로 생성하는 예제
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# K-food 관련 수집 자료 (요약 정리)
intro = (
    "K-food(한식)는 K-POP, K-드라마 등 한류 열풍과 함께 전 세계적으로 큰 인기를 끌고 있다. "
    "SNS와 유튜브를 통한 먹방(Mukbang) 콘텐츠 확산, 넷플릭스 드라마 속 음식 노출, "
    "간편하게 즐길 수 있는 K-편의식품의 해외 진출 등이 인기 요인으로 꼽힌다."
)

foods = [
    ("김치 (Kimchi)", "발효 배추김치로, 건강식품이자 한식의 대표 상징. 해외 대형마트에서도 쉽게 구매 가능."),
    ("불닭볶음면 (Buldak Fried Noodles)", "매운맛 챌린지로 유튜브·틱톡에서 화제가 되며 전 세계적으로 인기."),
    ("떡볶이 (Tteokbokki)", "쫄깃한 떡과 매콤달콤한 소스의 조합으로 길거리 음식의 대표주자."),
    ("한국식 치킨 (Korean Fried Chicken)", "이중 튀김 방식의 바삭한 식감과 다양한 양념(양념/간장/허니)이 특징."),
    ("비빔밥 (Bibimbap)", "다양한 나물과 고기를 고추장에 비벼 먹는 건강식으로 채식주의자에게도 인기."),
    ("삼겹살/K-BBQ", "직접 구워 먹는 스타일과 다양한 쌈 채소 조합이 해외에서 색다른 경험으로 인식."),
    ("김밥 (Gimbap)", "간편하게 즐길 수 있는 한 끼 식사로, 냉동김밥이 미국 등지에서 품절 대란을 일으킴."),
]

trends = [
    "미국 냉동김밥(Gimbap)이 대형 유통업체에서 완판 행렬을 기록.",
    "K-드라마·예능 속 먹방 장면이 해외 시청자의 한식 소비 욕구를 자극.",
    "글로벌 식품기업들이 K-소스(고추장, 양념치킨소스 등) 라인업을 확대.",
    "해외 대학가 및 대도시 중심으로 한식 프랜차이즈 매장 증가 추세.",
]

# 문서 생성
doc = Document()

title = doc.add_heading("K-Food, 요즘 왜 인기일까?", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("1. 개요", level=1)
doc.add_paragraph(intro)

doc.add_heading("2. 인기 K-Food 목록", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 1"
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "음식명"
hdr_cells[1].text = "설명"
for name, desc in foods:
    row_cells = table.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = desc

doc.add_heading("3. 최근 트렌드", level=1)
for t in trends:
    doc.add_paragraph(t, style="List Bullet")

doc.add_heading("4. 결론", level=1)
doc.add_paragraph(
    "K-food는 단순한 음식을 넘어 한류 콘텐츠와 결합된 문화 현상으로 확산되고 있으며, "
    "앞으로도 다양한 형태(밀키트, 소스, 간편식 등)로 글로벌 시장에서 성장할 것으로 전망된다."
)

# 저장 (python-docx는 .docx 형식만 지원)
doc.save("k-food.docx")
print("k-food.docx 파일이 생성되었습니다.")
