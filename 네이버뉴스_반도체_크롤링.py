import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


URL = "https://search.naver.com/search.naver"
PARAMS = {
    "where": "nexearch",
    "sm": "top_hty",
    "fbm": "0",
    "ie": "utf8",
    "query": "반도체",
}
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.0.0 Safari/537.36"
    )
}


def get_news_articles(url, params):
    response = requests.get(url, params=params, headers=HEADERS, timeout=10)
    response.raise_for_status()

    if "search.naver.com" not in response.url:
        raise RuntimeError(
            f"검색 페이지가 아닌 주소로 이동했습니다: {response.url}"
        )

    soup = BeautifulSoup(response.text, "html.parser")
    articles = []
    summaries_by_link = {}

    for summary_tag in soup.select('a[data-nlog-area$=".body"]'):
        link = summary_tag.get("href", "")
        if link:
            summaries_by_link[link] = summary_tag.get_text(" ", strip=True)

    for title_tag in soup.select('a[data-nlog-area$=".tit"]'):
        link = title_tag.get("href", "")
        articles.append(
            {
                "title": title_tag.get_text(" ", strip=True),
                "link": link,
                "summary": summaries_by_link.get(link, ""),
            }
        )

    return articles


articles = get_news_articles(URL, PARAMS)

workbook = Workbook()
worksheet = workbook.active
worksheet.title = "네이버 뉴스"

headers = ["제목", "링크", "요약"]
worksheet.append(headers)

for cell in worksheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(fill_type="solid", fgColor="2F75B5")

for article in articles:
    worksheet.append([article["title"], article["link"], article["summary"]])
    link_cell = worksheet.cell(worksheet.max_row, 2)
    link_cell.hyperlink = article["link"]
    link_cell.style = "Hyperlink"

worksheet.column_dimensions["A"].width = 55
worksheet.column_dimensions["B"].width = 80
worksheet.column_dimensions["C"].width = 100

for row in worksheet.iter_rows(min_row=2):
    row[0].alignment = row[1].alignment = row[2].alignment = (
        row[0].alignment.copy(wrap_text=True, vertical="top")
    )

workbook.save("naver_result.xlsx")

for number, article in enumerate(articles, start=1):
    print(f"[{number}] {article['title']}")
    print(f"링크: {article['link']}")
    print(f"요약: {article['summary']}")
    print("-" * 80)

print(f"총 {len(articles)}개의 뉴스 검색 결과를 naver_result.xlsx에 저장했습니다.")